
import json
import sys
from pathlib import Path

import requests
from docx import Document

MODEL = "gemma3"
OLLAMA_URL = "http://localhost:11434/api/chat"


def read_srs(file_path):
    """Extract paragraphs AND tables from a DOCX."""
    doc = Document(file_path)
    sections = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        sections.append(f"\n[TABLE {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            sections.append(" | ".join(cells))

    return "\n".join(sections)


def ask_model(system_prompt, user_prompt):
    response = requests.post(
        OLLAMA_URL,
        json={
            "model": MODEL,
            "stream": False,
            "format": "json",
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        },
        timeout=300,
    )
    response.raise_for_status()

    content = response.json()["message"]["content"]
    print("\nModel response preview:", content[:800])

    result = json.loads(content)

    if not isinstance(result, dict):
        raise ValueError("Model returned JSON, but not a JSON object.")

    return result


def generate_test_cases(srs_text):
    prompt = """
You are a software QA engineer generating developer test cases.

Read the supplied Software Requirements Specification (SRS).

Return a JSON object with this structure:
{
  "test_cases": [
    {
      "id": "TC-001",
      "requirement_reference": "requirement ID or Not specified",
      "scenario": "What is being tested",
      "preconditions": "Required setup",
      "test_steps": ["Step 1", "Step 2"],
      "test_data": "Input data",
      "expected_result": "Expected behavior",
      "test_type": "positive, negative, boundary, or error handling"
    }
  ],
  "open_questions": [],
  "requirements_needing_clarification": []
}

Rules:
- Generate test cases for the actual requirements in the SRS.
- Include positive, negative, and boundary tests when applicable.
- Use requirement IDs when present.
- Never invent unspecified business rules.
- Put unclear behavior in open_questions.
- Do not return an empty test_cases array if the SRS
  contains testable requirements.
- Return valid JSON only.
"""

    result = ask_model(
        prompt,
        "Generate developer test cases from this SRS:\n\n" + srs_text,
    )

    cases = result.get("test_cases", [])

    # Retry once if the first response has no usable test cases.
    if not isinstance(cases, list) or not cases:
        print("\nNo test cases generated. Retrying with a simpler prompt...")

        retry_prompt = """
Create developer test cases for the requirements below.
Return JSON with a test_cases array.
Each test case must have id, scenario, test_steps,
expected_result, and test_type.
Do not invent unspecified behavior.
Return at least one case per identifiable, testable requirement.
"""

        result = ask_model(retry_prompt, srs_text)
        cases = result.get("test_cases", [])

    if not isinstance(cases, list) or not cases:
        raise RuntimeError(
            "The model generated zero test cases after retry. "
            "Check the extracted SRS and model response above."
        )

    # Review is optional. Never allow a failed review to erase valid cases.
    print(f"\nGenerated {len(cases)} test cases. Reviewing...")

    try:
        reviewed = ask_model(
            """
Review these test cases against the original SRS.
Remove unsupported assumptions and duplicates.
Preserve valid test cases and their expected results.
Do not return an empty test_cases array if valid cases exist.
Return JSON with test_cases, open_questions, and
requirements_needing_clarification.
""",
            "SRS:\n" + srs_text
            + "\n\nGenerated test cases:\n"
            + json.dumps(result, ensure_ascii=False),
        )

        reviewed_cases = reviewed.get("test_cases", [])

        if isinstance(reviewed_cases, list) and reviewed_cases:
            result = reviewed
        else:
            print("Review returned no cases. Keeping original results.")

    except (requests.RequestException, ValueError, KeyError) as exc:
        print(f"Review failed; keeping original results: {exc}")

    return result


def main():
    if len(sys.argv) != 2:
        print("Usage: python srs_test_agent.py requirements.docx")
        return

    file_path = Path(sys.argv[1])

    if not file_path.is_file():
        raise FileNotFoundError(f"SRS file not found: {file_path}")

    srs_text = read_srs(str(file_path))

    print("Extracted SRS characters:", len(srs_text))
    print("Extracted SRS preview:\n", srs_text[:1000])

    if not srs_text.strip():
        raise ValueError(
            "No text extracted. Check whether the DOCX contains "
            "images/scanned pages or content inside text boxes."
        )

    result = generate_test_cases(srs_text)

    output_path = file_path.with_name(
        file_path.stem + "_test_cases.json"
    )

    output_path.write_text(
        json.dumps(result, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print(f"\nSaved: {output_path.resolve()}")
    print("Final test case count:", len(result["test_cases"]))


if __name__ == "__main__":
    main()
