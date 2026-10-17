import re
from pathlib import Path

from docx import Document

from config import (
    MAX_CHUNK_CHARACTERS,
    CHUNK_OVERLAP_BLOCKS
)


HEADING_PATTERNS = [
    # Examples:
    # 15480_FR_15: 800-RTT Service Scalability
    # PROJECT_NFR_10: Performance
    re.compile(
        r"^(?P<section_id>[A-Za-z0-9-]+_"
        r"(?:FR|IR|AR|NFR|REQ)_\d+)"
        r"\s*:\s*(?P<title>.+)$",
        re.IGNORECASE
    ),

    # Examples:
    # FR_15: Scalability
    # NFR_3: Performance
    re.compile(
        r"^(?P<section_id>"
        r"(?:FR|IR|AR|NFR|REQ)_\d+)"
        r"\s*:\s*(?P<title>.+)$",
        re.IGNORECASE
    ),

    # Examples:
    # 3.4.1 Service Scalability
    # 5.2 Architecture
    re.compile(
        r"^(?P<section_id>\d+(?:\.\d+)+)"
        r"\s+(?P<title>.+)$"
    ),

    # Examples:
    # Functional Requirement 12: Scalability
    # Non-Functional Requirement 4: Security
    re.compile(
        r"^(?P<section_id>"
        r"(?:Functional|Integration|Architecture|"
        r"Non-Functional)\s+Requirement\s+\d+)"
        r"\s*:\s*(?P<title>.+)$",
        re.IGNORECASE
    )
]


def normalize_text(text):
    return re.sub(
        r"\s+",
        " ",
        text
    ).strip()


def detect_heading(text):
    """
    Detect a structured heading from plain text.

    Returns:
        {
            "section_id": "...",
            "section_title": "..."
        }

    Returns None when the block is not a heading.
    """

    cleaned = text.strip()

    if cleaned.lower().startswith("section:"):
        cleaned = cleaned[len("section:"):].strip()

    for pattern in HEADING_PATTERNS:
        match = pattern.match(cleaned)

        if match:
            return {
                "section_id": match.group(
                    "section_id"
                ).strip(),
                "section_title": match.group(
                    "title"
                ).strip()
            }

    return None


def extract_docx_blocks(file_path):
    """
    Extract paragraphs and table rows from a DOCX.

    Each extracted item is represented as a block.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Document not found: {file_path.resolve()}"
        )

    document = Document(str(file_path))
    blocks = []

    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)

        if not text:
            continue

        style_name = ""

        if paragraph.style is not None:
            style_name = (
                paragraph.style.name or ""
            )

        is_heading_style = (
            style_name.lower().startswith("heading")
        )

        blocks.append({
            "text": text,
            "is_heading_style": is_heading_style,
            "source_type": "paragraph"
        })

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):
        for row_index, row in enumerate(
            table.rows,
            start=1
        ):
            cells = []

            for cell in row.cells:
                cell_text = normalize_text(
                    cell.text
                )

                if cell_text:
                    cells.append(cell_text)

            if not cells:
                continue

            blocks.append({
                "text": (
                    f"TABLE {table_index}, "
                    f"ROW {row_index}: "
                    + " | ".join(cells)
                ),
                "is_heading_style": False,
                "source_type": "table"
            })

    return blocks


def build_sections(blocks):
    """
    Group blocks under their nearest section heading.

    This avoids mixing unrelated functional requirements
    into arbitrary character-sized chunks.
    """

    sections = []

    current_section = {
        "section_id": "GENERAL",
        "section_title": "General",
        "blocks": []
    }

    for block in blocks:
        text = block["text"]

        detected_heading = detect_heading(text)

        if (
            detected_heading is None
            and block["is_heading_style"]
        ):
            detected_heading = {
                "section_id": "HEADING",
                "section_title": text
            }

        if detected_heading is not None:
            if current_section["blocks"]:
                sections.append(current_section)

            current_section = {
                "section_id": detected_heading[
                    "section_id"
                ],
                "section_title": detected_heading[
                    "section_title"
                ],
                "blocks": [text]
            }

        else:
            current_section["blocks"].append(text)

    if current_section["blocks"]:
        sections.append(current_section)

    return sections


def split_section(section):
    """
    Split only oversized sections.

    The section ID and title are repeated in every child
    chunk so that no child loses its structural context.
    """

    section_id = section["section_id"]
    section_title = section["section_title"]

    header = (
        f"SECTION: {section_id}: {section_title}"
    )

    content_blocks = list(section["blocks"])

    if (
        content_blocks
        and content_blocks[0].lower().startswith(
            "section:"
        )
    ):
        content_blocks = content_blocks[1:]

    chunks = []
    current_blocks = [header]
    current_length = len(header)

    for block in content_blocks:
        additional_length = len(block) + 1

        if (
            len(current_blocks) > 1
            and current_length + additional_length
            > MAX_CHUNK_CHARACTERS
        ):
            chunks.append(
                "\n".join(current_blocks)
            )

            overlap = current_blocks[
                -CHUNK_OVERLAP_BLOCKS:
            ]

            overlap = [
                value
                for value in overlap
                if value != header
            ]

            current_blocks = [header] + overlap

            current_length = sum(
                len(value) + 1
                for value in current_blocks
            )

        current_blocks.append(block)
        current_length += additional_length

    if len(current_blocks) > 1:
        chunks.append(
            "\n".join(current_blocks)
        )

    elif not chunks:
        chunks.append(header)

    return chunks


def read_docx_as_chunks(file_path):
    """
    Return chunks and associated metadata.
    """

    file_path = Path(file_path)

    blocks = extract_docx_blocks(file_path)
    sections = build_sections(blocks)

    chunks = []
    metadatas = []

    document_type = classify_document(
        file_path.name
    )

    for section_number, section in enumerate(
        sections
    ):
        section_chunks = split_section(section)

        for child_index, chunk in enumerate(
            section_chunks
        ):
            chunks.append(chunk)

            metadatas.append({
                "source": file_path.name,
                "document_type": document_type,
                "section_id": section[
                    "section_id"
                ],
                "section_title": section[
                    "section_title"
                ],
                "section_number": section_number,
                "child_index": child_index
            })

    return chunks, metadatas


def classify_document(filename):
    """
    Basic, generic document classification.

    This can later be replaced by metadata from
    SharePoint, ADO, QC or an LLM classifier.
    """

    lowered_name = filename.lower()

    if "srs" in lowered_name:
        return "SRS"

    if "rca" in lowered_name:
        return "RCA"

    if (
        "incident" in lowered_name
        or "production" in lowered_name
    ):
        return "INCIDENT"

    if "design" in lowered_name:
        return "DESIGN"

    if (
        "test" in lowered_name
        or "validation" in lowered_name
    ):
        return "TEST"

    if "requirement" in lowered_name:
        return "REQUIREMENT"

    return "GENERAL"
