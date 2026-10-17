import hashlib
import re
from pathlib import Path

import chromadb
import ollama
from docx import Document


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIRECTORY = Path(__file__).resolve().parent

DOCUMENT_DIRECTORY = BASE_DIRECTORY / "documents"
CHROMA_DIRECTORY = BASE_DIRECTORY / "chroma_db"

COLLECTION_NAME = "engineering_knowledge"

EMBEDDING_MODEL = "nomic-embed-text"

# Maximum approximate size of one chunk.
MAX_CHUNK_CHARACTERS = 2500

# Number of previous content blocks repeated when an oversized
# section is split.
CHUNK_OVERLAP_BLOCKS = 1

# Number of chunks sent to Ollama in a single embedding request.
EMBEDDING_BATCH_SIZE = 16

# Number of records inserted into ChromaDB in one operation.
CHROMA_BATCH_SIZE = 100


# ============================================================
# GENERIC HEADING PATTERNS
# ============================================================

HEADING_PATTERNS = [
    # Examples:
    # 15480_FR_15: 800-RTT Service Scalability
    # ABC_NFR_10: Performance
    re.compile(
        r"^(?P<section_id>"
        r"[A-Za-z0-9-]+_(?:FR|IR|AR|NFR|REQ)_\d+"
        r")\s*:\s*(?P<section_title>.+)$",
        re.IGNORECASE
    ),

    # Examples:
    # FR_15: Scalability
    # NFR_3: Performance
    re.compile(
        r"^(?P<section_id>"
        r"(?:FR|IR|AR|NFR|REQ)_\d+"
        r")\s*:\s*(?P<section_title>.+)$",
        re.IGNORECASE
    ),

    # Examples:
    # 3.4.1 Service Scalability
    # 5.2 Deployment Architecture
    re.compile(
        r"^(?P<section_id>"
        r"\d+(?:\.\d+)+"
        r")\s+(?P<section_title>.+)$"
    ),

    # Examples:
    # Functional Requirement 12: Scalability
    # Non-Functional Requirement 5: Security
    re.compile(
        r"^(?P<section_id>"
        r"(?:Functional|Integration|Architecture|"
        r"Non-Functional)\s+Requirement\s+\d+"
        r")\s*:\s*(?P<section_title>.+)$",
        re.IGNORECASE
    )
]


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text):
    """
    Normalize repeated spaces and line breaks.

    This keeps extracted content clean without changing the
    actual words in the document.
    """

    return re.sub(
        r"\s+",
        " ",
        text
    ).strip()


# ============================================================
# DOCUMENT CLASSIFICATION
# ============================================================

def classify_document(filename):
    """
    Perform basic document classification from the filename.

    This is metadata only. It does not change document content
    or retrieval behavior.
    """

    lowered_name = filename.lower()

    if "srs" in lowered_name:
        return "SRS"

    if "rca" in lowered_name:
        return "RCA"

    if (
        "incident" in lowered_name
        or "production" in lowered_name
    ):
        return "INCIDENT"

    if "design" in lowered_name:
        return "DESIGN"

    if (
        "test" in lowered_name
        or "validation" in lowered_name
    ):
        return "TEST"

    if "requirement" in lowered_name:
        return "REQUIREMENT"

    if "troubleshoot" in lowered_name:
        return "TROUBLESHOOTING"

    return "GENERAL"


# ============================================================
# HEADING DETECTION
# ============================================================

def detect_heading(text):
    """
    Detect structured SRS and numbered headings.

    Returns a dictionary when the supplied text is recognized
    as a heading.

    Otherwise, returns None.
    """

    cleaned_text = text.strip()

    if cleaned_text.lower().startswith("section:"):
        cleaned_text = cleaned_text[
            len("section:"):
        ].strip()

    for pattern in HEADING_PATTERNS:
        match = pattern.match(cleaned_text)

        if match:
            return {
                "section_id": match.group(
                    "section_id"
                ).strip(),
                "section_title": match.group(
                    "section_title"
                ).strip()
            }

    return None


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_blocks(file_path):
    """
    Extract normal paragraphs and table rows from a DOCX file.

    Each extracted element becomes a block.

    Word heading styles are preserved through the
    is_heading_style property.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Document does not exist: "
            f"{file_path.resolve()}"
        )

    document = Document(str(file_path))

    blocks = []

    paragraph_count = 0
    table_row_count = 0

    # Extract ordinary paragraphs.
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)

        if not text:
            continue

        paragraph_count += 1

        style_name = ""

        if paragraph.style is not None:
            style_name = (
                paragraph.style.name or ""
            )

        is_heading_style = (
            style_name.lower().startswith("heading")
        )

        blocks.append({
            "text": text,
            "source_type": "paragraph",
            "is_heading_style": is_heading_style,
            "style_name": style_name
        })

    # Extract tables.
    for table_index, table in enumerate(
        document.tables,
        start=1
    ):
        for row_index, row in enumerate(
            table.rows,
            start=1
        ):
            cells = []

            for cell in row.cells:
                cell_text = normalize_text(
                    cell.text
                )

                if cell_text:
                    cells.append(cell_text)

            if not cells:
                continue

            table_row_count += 1

            row_text = " | ".join(cells)

            blocks.append({
                "text": (
                    f"TABLE {table_index}, "
                    f"ROW {row_index}: "
                    f"{row_text}"
                ),
                "source_type": "table",
                "is_heading_style": False,
                "style_name": ""
            })

    print(
        f"  Paragraphs extracted: "
        f"{paragraph_count}"
    )

    print(
        f"  Table rows extracted: "
        f"{table_row_count}"
    )

    print(
        f"  Total extracted blocks: "
        f"{len(blocks)}"
    )

    return blocks


# ============================================================
# SECTION BUILDING
# ============================================================

def build_sections(blocks):
    """
    Group document blocks under section headings.

    A section continues until the next detected heading.

    This avoids arbitrary character chunks mixing unrelated
    functional requirements.
    """

    sections = []

    current_section = {
        "section_id": "GENERAL",
        "section_title": "General",
        "blocks": []
    }

    generic_heading_counter = 0

    for block in blocks:
        text = block["text"]

        detected_heading = detect_heading(text)

        # If the paragraph uses a Word Heading style but does
        # not match an SRS requirement pattern, preserve it as
        # a generic section.
        if (
            detected_heading is None
            and block["is_heading_style"]
        ):
            generic_heading_counter += 1

            detected_heading = {
                "section_id": (
                    f"HEADING_{generic_heading_counter}"
                ),
                "section_title": text
            }

        if detected_heading is not None:
            # Save the previous section first.
            if current_section["blocks"]:
                sections.append(current_section)

            current_section = {
                "section_id": detected_heading[
                    "section_id"
                ],
                "section_title": detected_heading[
                    "section_title"
                ],
                "blocks": []
            }

        else:
            current_section["blocks"].append(text)

    if current_section["blocks"]:
        sections.append(current_section)

    return sections


# ============================================================
# SECTION SPLITTING
# ============================================================

def split_section(section):
    """
    Split a single section only if it exceeds the configured
    maximum chunk size.

    Every child chunk repeats the section ID and title so the
    structural context is never lost.
    """

    section_id = section["section_id"]
    section_title = section["section_title"]

    header = (
        f"SECTION: {section_id}: "
        f"{section_title}"
    )

    content_blocks = section["blocks"]

    # Handle an empty heading section.
    if not content_blocks:
        return [header]

    chunks = []

    current_blocks = [header]
    current_length = len(header)

    for content_block in content_blocks:
        block_length = len(content_block) + 1

        should_close_chunk = (
            len(current_blocks) > 1
            and current_length + block_length
            > MAX_CHUNK_CHARACTERS
        )

        if should_close_chunk:
            chunks.append(
                "\n".join(current_blocks)
            )

            # Preserve a small amount of context in the next
            # child chunk.
            overlap_blocks = current_blocks[
                -CHUNK_OVERLAP_BLOCKS:
            ]

            overlap_blocks = [
                block
                for block in overlap_blocks
                if block != header
            ]

            current_blocks = [
                header,
                *overlap_blocks
            ]

            current_length = sum(
                len(block) + 1
                for block in current_blocks
            )

        current_blocks.append(content_block)
        current_length += block_length

    if len(current_blocks) > 1:
        chunks.append(
            "\n".join(current_blocks)
        )

    if not chunks:
        chunks.append(header)

    return chunks


# ============================================================
# CONVERT ONE DOCX INTO CHUNKS
# ============================================================

def read_docx_as_chunks(file_path):
    """
    Convert one DOCX document into searchable chunks and
    metadata dictionaries.
    """

    file_path = Path(file_path)

    blocks = extract_docx_blocks(file_path)

    if not blocks:
        print(
            "  WARNING: No text blocks were extracted."
        )

        return [], []

    sections = build_sections(blocks)

    print(
        f"  Sections detected: {len(sections)}"
    )

    document_type = classify_document(
        file_path.name
    )

    chunks = []
    metadatas = []

    for section_number, section in enumerate(
        sections,
        start=1
    ):
        section_chunks = split_section(section)

        for child_index, chunk in enumerate(
            section_chunks
        ):
            chunks.append(chunk)

            metadatas.append({
                "source": file_path.name,
                "document_type": document_type,
                "section_id": section[
                    "section_id"
                ],
                "section_title": section[
                    "section_title"
                ],
                "section_number": section_number,
                "child_index": child_index
            })

    print(
        f"  Searchable chunks created: "
        f"{len(chunks)}"
    )

    return chunks, metadatas


# ============================================================
# DOCUMENT DISCOVERY
# ============================================================

def find_docx_files(directory):
    """
    Recursively find DOCX files.

    Temporary Microsoft Word files beginning with ~$ are
    ignored.
    """

    directory = Path(directory)

    directory.mkdir(
        parents=True,
        exist_ok=True
    )

    files = []

    for file_path in directory.rglob("*.docx"):
        if file_path.name.startswith("~$"):
            continue

        files.append(file_path)

    return sorted(files)


# ============================================================
# STABLE CHUNK IDENTIFIERS
# ============================================================

def create_chunk_id(
    source,
    section_id,
    child_index,
    content
):
    """
    Create a deterministic SHA-256 identifier for a chunk.

    The same chunk will receive the same identifier on the
    next ingestion run.
    """

    identity = (
        f"{source}|"
        f"{section_id}|"
        f"{child_index}|"
        f"{content}"
    )

    return hashlib.sha256(
        identity.encode("utf-8")
    ).hexdigest()


# ============================================================
# OLLAMA VALIDATION
# ============================================================

def validate_ollama_and_model():
    """
    Validate the connection to Ollama and confirm that the
    configured embedding model is installed.
    """

    try:
        response = ollama.list()

    except Exception as exception:
        raise RuntimeError(
            "Could not connect to Ollama.\n"
            "Ensure the Ollama application is running.\n"
            "Test with:\n"
            "ollama list"
        ) from exception

    installed_models = set()

    for model in response.models:
        model_name = model.model

        installed_models.add(model_name)

        installed_models.add(
            model_name.split(":")[0]
        )

    configured_model_base = (
        EMBEDDING_MODEL.split(":")[0]
    )

    model_found = (
        EMBEDDING_MODEL in installed_models
        or configured_model_base
        in installed_models
    )

    if not model_found:
        raise RuntimeError(
            f'The embedding model '
            f'"{EMBEDDING_MODEL}" is not installed.\n\n'
            f"Run:\n"
            f"ollama pull {EMBEDDING_MODEL}"
        )


# ============================================================
# EMBEDDING GENERATION
# ============================================================

def create_embeddings(texts):
    """
    Generate embeddings for one batch of text chunks.
    """

    if not texts:
        return []

    response = ollama.embed(
        model=EMBEDDING_MODEL,
        input=texts
    )

    embeddings = response["embeddings"]

    if len(embeddings) != len(texts):
        raise RuntimeError(
            "Ollama returned an unexpected embedding count.\n"
            f"Text chunks: {len(texts)}\n"
            f"Embeddings: {len(embeddings)}"
        )

    return embeddings


def create_embeddings_in_batches(chunks):
    """
    Generate embeddings in manageable batches.
    """

    all_embeddings = []
    total_chunks = len(chunks)

    print()
    print("=" * 70)
    print("GENERATING EMBEDDINGS")
    print("=" * 70)

    print(
        f"Embedding model: {EMBEDDING_MODEL}"
    )

    print(
        f"Batch size: {EMBEDDING_BATCH_SIZE}"
    )

    for start_index in range(
        0,
        total_chunks,
        EMBEDDING_BATCH_SIZE
    ):
        end_index = min(
            start_index + EMBEDDING_BATCH_SIZE,
            total_chunks
        )

        batch = chunks[
            start_index:end_index
        ]

        print(
            f"Embedding chunks "
            f"{start_index + 1}-"
            f"{end_index} "
            f"of {total_chunks}"
        )

        try:
            batch_embeddings = create_embeddings(
                batch
            )

        except Exception as exception:
            raise RuntimeError(
                f"Embedding generation failed for "
                f"chunks {start_index + 1}-"
                f"{end_index}."
            ) from exception

        all_embeddings.extend(
            batch_embeddings
        )

    if len(all_embeddings) != total_chunks:
        raise RuntimeError(
            "The final embedding count does not match "
            "the chunk count."
        )

    return all_embeddings


# ============================================================
# COLLECT CHUNKS FROM ALL DOCUMENTS
# ============================================================

def collect_all_chunks(docx_files):
    """
    Read every discovered DOCX file and combine the generated
    chunks, IDs and metadata.
    """

    all_ids = []
    all_chunks = []
    all_metadatas = []

    successful_documents = 0
    failed_documents = 0

    for document_index, file_path in enumerate(
        docx_files,
        start=1
    ):
        print()
        print("=" * 70)

        print(
            f"DOCUMENT "
            f"{document_index}/{len(docx_files)}"
        )

        print(file_path)
        print("=" * 70)

        try:
            chunks, metadatas = read_docx_as_chunks(
                file_path
            )

        except Exception as exception:
            failed_documents += 1

            print(
                f"ERROR: Failed to process "
                f"{file_path.name}"
            )

            print(exception)
            print("The document will be skipped.")

            continue

        if len(chunks) != len(metadatas):
            failed_documents += 1

            print(
                "ERROR: Chunk count and metadata count "
                "do not match."
            )

            print(
                f"Chunks: {len(chunks)}"
            )

            print(
                f"Metadata records: "
                f"{len(metadatas)}"
            )

            continue

        if not chunks:
            failed_documents += 1

            print(
                "WARNING: Document produced no chunks."
            )

            continue

        successful_documents += 1

        try:
            relative_source = file_path.relative_to(
                DOCUMENT_DIRECTORY
            ).as_posix()

        except ValueError:
            relative_source = file_path.name

        for chunk, metadata in zip(
            chunks,
            metadatas
        ):
            if not chunk.strip():
                continue

            copied_metadata = dict(metadata)

            copied_metadata[
                "source"
            ] = relative_source

            # Chroma metadata values must remain primitive
            # strings, integers, floats or booleans.
            copied_metadata[
                "absolute_path"
            ] = str(file_path.resolve())

            section_id = copied_metadata.get(
                "section_id",
                "GENERAL"
            )

            child_index = int(
                copied_metadata.get(
                    "child_index",
                    0
                )
            )

            chunk_id = create_chunk_id(
                source=relative_source,
                section_id=section_id,
                child_index=child_index,
                content=chunk
            )

            all_ids.append(chunk_id)
            all_chunks.append(chunk)
            all_metadatas.append(
                copied_metadata
            )

    return {
        "ids": all_ids,
        "chunks": all_chunks,
        "metadatas": all_metadatas,
        "successful_documents": successful_documents,
        "failed_documents": failed_documents
    }


# ============================================================
# CHROMADB REBUILD
# ============================================================

def rebuild_chroma_collection(
    ids,
    chunks,
    metadatas,
    embeddings
):
    """
    Completely rebuild the ChromaDB collection.

    Rebuilding is intentional during POC development. It
    prevents old character-based chunks from remaining beside
    new section-based chunks.
    """

    CHROMA_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True
    )

    client = chromadb.PersistentClient(
        path=str(CHROMA_DIRECTORY)
    )

    collection_names = {
        collection.name
        for collection in client.list_collections()
    }

    if COLLECTION_NAME in collection_names:
        print()
        print(
            f'Deleting existing collection '
            f'"{COLLECTION_NAME}"...'
        )

        client.delete_collection(
            name=COLLECTION_NAME
        )

    print(
        f'Creating collection '
        f'"{COLLECTION_NAME}"...'
    )

    collection = client.create_collection(
        name=COLLECTION_NAME,
        metadata={
            "hnsw:space": "cosine"
        }
    )

    total_chunks = len(chunks)

    print()
    print("=" * 70)
    print("STORING RECORDS IN CHROMADB")
    print("=" * 70)

    for start_index in range(
        0,
        total_chunks,
        CHROMA_BATCH_SIZE
    ):
        end_index = min(
            start_index + CHROMA_BATCH_SIZE,
            total_chunks
        )

        collection.add(
            ids=ids[start_index:end_index],
            embeddings=embeddings[
                start_index:end_index
            ],
            documents=chunks[
                start_index:end_index
            ],
            metadatas=metadatas[
                start_index:end_index
            ]
        )

        print(
            f"Stored records "
            f"{start_index + 1}-"
            f"{end_index} "
            f"of {total_chunks}"
        )

    return collection


# ============================================================
# DEBUG SUMMARY
# ============================================================

def print_section_summary(metadatas):
    """
    Print detected sections so chunking problems can be seen
    before retrieval is tested.
    """

    seen_sections = set()

    print()
    print("=" * 70)
    print("DETECTED SECTIONS")
    print("=" * 70)

    for metadata in metadatas:
        section_key = (
            metadata.get("source"),
            metadata.get("section_id"),
            metadata.get("section_title")
        )

        if section_key in seen_sections:
            continue

        seen_sections.add(section_key)

        print(
            f"{metadata.get('source')} | "
            f"{metadata.get('section_id')} | "
            f"{metadata.get('section_title')}"
        )


def print_ingestion_summary(
    discovered_documents,
    ingestion_result,
    collection
):
    """
    Print final ingestion statistics.
    """

    print()
    print("=" * 70)
    print("INGESTION COMPLETED")
    print("=" * 70)

    print(
        f"Documents discovered: "
        f"{len(discovered_documents)}"
    )

    print(
        f"Documents processed successfully: "
        f"{ingestion_result['successful_documents']}"
    )

    print(
        f"Documents failed: "
        f"{ingestion_result['failed_documents']}"
    )

    print(
        f"Chunks generated: "
        f"{len(ingestion_result['chunks'])}"
    )

    print(
        f"Records stored in ChromaDB: "
        f"{collection.count()}"
    )

    print(
        f"Collection name: "
        f"{COLLECTION_NAME}"
    )

    print(
        f"ChromaDB location: "
        f"{CHROMA_DIRECTORY}"
    )

    document_type_counts = {}

    for metadata in ingestion_result[
        "metadatas"
    ]:
        document_type = metadata.get(
            "document_type",
            "GENERAL"
        )

        current_count = document_type_counts.get(
            document_type,
            0
        )

        document_type_counts[
            document_type
        ] = current_count + 1

    print()
    print("Chunks by document type:")

    for document_type, count in sorted(
        document_type_counts.items()
    ):
        print(
            f"  {document_type}: {count}"
        )


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 70)
    print("ENGINEERING KNOWLEDGE INGESTION")
    print("=" * 70)

    print(
        f"Document directory: "
        f"{DOCUMENT_DIRECTORY}"
    )

    print(
        f"ChromaDB directory: "
        f"{CHROMA_DIRECTORY}"
    )

    print(
        f"Collection name: "
        f"{COLLECTION_NAME}"
    )

    print(
        f"Embedding model: "
        f"{EMBEDDING_MODEL}"
    )

    print()
    print("Checking Ollama and embedding model...")

    validate_ollama_and_model()

    print(
        "Ollama connection and embedding model "
        "are available."
    )

    docx_files = find_docx_files(
        DOCUMENT_DIRECTORY
    )

    if not docx_files:
        raise RuntimeError(
            f"No DOCX files were found under:\n"
            f"{DOCUMENT_DIRECTORY}\n\n"
            f"Create the documents directory and copy "
            f"one or more .docx files into it."
        )

    print()
    print(
        f"DOCX files discovered: "
        f"{len(docx_files)}"
    )

    for file_path in docx_files:
        print(
            f"  {file_path.relative_to(DOCUMENT_DIRECTORY)}"
        )

    ingestion_result = collect_all_chunks(
        docx_files
    )

    ids = ingestion_result["ids"]
    chunks = ingestion_result["chunks"]
    metadatas = ingestion_result["metadatas"]

    if not chunks:
        raise RuntimeError(
            "Documents were discovered, but no searchable "
            "chunks were created."
        )

    if not (
        len(ids)
        == len(chunks)
        == len(metadatas)
    ):
        raise RuntimeError(
            "ID, chunk and metadata counts are different.\n"
            f"IDs: {len(ids)}\n"
            f"Chunks: {len(chunks)}\n"
            f"Metadata records: {len(metadatas)}"
        )

    print_section_summary(metadatas)

    print()
    print(
        f"Total searchable chunks: "
        f"{len(chunks)}"
    )

    embeddings = create_embeddings_in_batches(
        chunks
    )

    collection = rebuild_chroma_collection(
        ids=ids,
        chunks=chunks,
        metadatas=metadatas,
        embeddings=embeddings
    )

    print_ingestion_summary(
        discovered_documents=docx_files,
        ingestion_result=ingestion_result,
        collection=collection
    )


if __name__ == "__main__":
    main()